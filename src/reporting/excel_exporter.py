"""
excel_exporter.py
==================
Writes the *raw* per-model evaluation results to an Excel workbook: a summary
sheet of scalar metrics, a per-class classification-report sheet, and a
confusion-matrix sheet.

Only takes metrics_merged (from metrics.merge()). metrics_merged_ranked is
intentionally never referenced here -- Excel is meant to stay a plain, literal
dump of the numbers for tinkering, not an opinionated "best/worst" view. That
comparison lives in chart_builder.py / report_builder.py instead, which feed the
PDF/PNG report.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import reporting.utils as UTILS
import config

def export(metrics_merged, path="report.xlsx", checkpoints=None):
    model_ids = list(metrics_merged["model_id"])
    keys = UTILS.scalar_keys(metrics_merged)

    wb = Workbook()
    _write_summary_sheet(wb.active, metrics_merged, model_ids, keys, checkpoints)
    wb.active.title = "Summary"
    _write_class_report_sheet(wb.create_sheet("Per-Class Report"), metrics_merged, model_ids)
    _write_confusion_sheet(wb.create_sheet("Confusion Matrices"), metrics_merged, model_ids)
    wb.save(path)
    return path


def _style_header(ws, row, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=UTILS.FONT_NAME, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=UTILS.HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _display_len(value):
    # width of what Excel shows, not of the stored value: floats are shown with VALUE_FORMAT
    return len(f"{value:.4f}") if isinstance(value, float) else len(str(value))


def _autofit(ws, n_cols, min_width=10, max_width=32):
    for c in range(1, n_cols + 1):
        col = get_column_letter(c)
        longest = max(
            (_display_len(cell.value) for cell in ws[col] if cell.value is not None),
            default=min_width,
        )
        ws.column_dimensions[col].width = max(min_width, min(max_width, longest + 2))


def _write_number(ws, row, column, value, number_format):
    cell = ws.cell(row=row, column=column, value=value)
    cell.number_format = number_format
    cell.font = Font(name=UTILS.FONT_NAME)
    cell.alignment = Alignment(horizontal="center")


def _write_summary_sheet(ws, metrics_merged, model_ids, keys, checkpoints):
    headers = ["Model"] + [UTILS.METRIC_LABELS[k] for k in keys]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)

    for i, model_id in enumerate(model_ids):
        row = i + 2
        ws.cell(row=row, column=1, value=model_id).font = Font(name=UTILS.FONT_NAME, bold=True)
        for j, key in enumerate(keys):
            value = metrics_merged[key][i]
            _write_number(ws, row, j + 2, "N/A" if value != value else float(value), UTILS.VALUE_FORMAT)

    if checkpoints:
        first_col = len(headers) + 1
        for k, (field, header, number_format) in enumerate(UTILS.CHECKPOINT_COLUMNS):
            ws.cell(row=1, column=first_col + k, value=header)
            for i, checkpoint in enumerate(checkpoints):
                _write_number(ws, i + 2, first_col + k, checkpoint[field], number_format)

    _style_header(ws, 1, ws.max_column)
    ws.freeze_panes = "B2"
    _autofit(ws, ws.max_column)


def _write_class_report_sheet(ws, metrics_merged, model_ids):
    headers = ["Model", "Class", "Precision", "Recall", "F1-score", "Support"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    _style_header(ws, 1, len(headers))

    reports = metrics_merged["classification_report"]
    row = 2
    for i, model_id in enumerate(model_ids):
        cr = reports[i]
        rows_for_model = UTILS.extract_class_names(cr) + ["macro avg", "weighted avg"]
        for cls in rows_for_model:
            entry = cr[cls]
            ws.cell(row=row, column=1, value=model_id).font = Font(name=UTILS.FONT_NAME)
            ws.cell(row=row, column=2, value=cls).font = Font(name=UTILS.FONT_NAME)
            for col, (stat, number_format) in enumerate(UTILS.CLASS_REPORT_COLUMNS, start=3):
                _write_number(ws, row, col, entry[stat], number_format)
            row += 1
        row += 1  # blank row between models
    _autofit(ws, len(headers))


def _write_confusion_sheet(ws, metrics_merged, model_ids):
    matrices = metrics_merged["confusion_matrix"]
    class_names = UTILS.extract_class_names(metrics_merged["classification_report"][0])
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    row_cursor = 1
    for i, model_id in enumerate(model_ids):
        ws.cell(row=row_cursor, column=1, value=f"{model_id} — rows: actual, columns: predicted").font = Font(name=UTILS.FONT_NAME, bold=True)
        row_cursor += 1

        header_row = row_cursor
        ws.cell(row=header_row, column=1, value="")
        for j, cname in enumerate(class_names):
            ws.cell(row=header_row, column=j + 2, value=cname)
        _style_header(ws, header_row, len(class_names) + 1)

        cm = matrices[i]
        for r, cname in enumerate(class_names):
            data_row = header_row + 1 + r
            ws.cell(row=data_row, column=1, value=cname).font = Font(name=UTILS.FONT_NAME, bold=True)
            for c in range(len(class_names)):
                cell = ws.cell(row=data_row, column=c + 2, value=int(cm[r, c]))
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
                if r == c:
                    cell.fill = PatternFill("solid", fgColor=UTILS.DIAGONAL_FILL)
        row_cursor = header_row + 1 + len(class_names) + 2

    _autofit(ws, len(class_names) + 1)
